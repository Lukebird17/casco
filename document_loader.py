import os
import subprocess
import tempfile
import shutil
import json
import hashlib
from typing import List, Dict, Optional, Tuple
from pathlib import Path

from config import DATA_DIR, LIBREOFFICE_PATH
from enhanced_ocr import EnhancedOCRProcessor, MINERU_AVAILABLE
from image_describer import ImageDescriber

# 中间数据目录（用于缓存）
INTERMEDIATE_DIR = os.path.join(os.path.dirname(DATA_DIR), "intermediate_data")


class DocumentLoader:
    def __init__(
        self,
        data_dir: str = DATA_DIR,
        enable_image_description: bool = True,
        save_intermediate_dir: str = INTERMEDIATE_DIR,
    ):
        self.data_dir = data_dir
        self.save_intermediate_dir = save_intermediate_dir
        self.supported_formats = [".pdf", ".pptx", ".docx", ".txt", ".md"]
        self.enable_image_description = enable_image_description
        
        # 确保缓存目录存在
        if self.save_intermediate_dir and not os.path.exists(self.save_intermediate_dir):
            os.makedirs(self.save_intermediate_dir, exist_ok=True)
            print(f"✅ 创建缓存目录: {self.save_intermediate_dir}")
        
        # 初始化MinerU处理器
        if MINERU_AVAILABLE:
            self.ocr_processor = EnhancedOCRProcessor()
            print("✅ MinerU已启用，将用于处理PDF")
        else:
            self.ocr_processor = None
            print("⚠️  MinerU不可用，将使用基础解析器")
        
        # 初始化图片描述生成器
        if self.enable_image_description:
            try:
                self.image_describer = ImageDescriber()
                print("✅ 图片描述生成器已启用")
            except Exception as e:
                print(f"⚠️  图片描述生成器初始化失败: {e}")
                self.image_describer = None
        else:
            self.image_describer = None
    
    def _add_image_descriptions_and_embed(
        self, 
        images: List[Dict], 
        pages: List[Dict]
    ) -> List[Dict]:
        """
        为图片生成描述并将描述嵌入到对应页面的文本内容中
        
        Args:
            images: 图片信息列表
            pages: 页面文本列表（会被修改）
            
        Returns:
            更新后的图片列表（添加了description字段）
        """
        print(f"\n📝 开始为 {len(images)} 张图片生成描述...")
        
        # 批量生成图片描述（最多并发3个）
        images = self.image_describer.batch_describe_images(images, max_workers=3)
        
        # 按页码对图片分组
        page_to_images = {}
        for img in images:
            page_num = img.get('page_number', 1)
            if page_num not in page_to_images:
                page_to_images[page_num] = []
            page_to_images[page_num].append(img)
        
        # 将图片描述插入到对应页面的文本中
        print(f"\n📋 将图片描述插入到文本内容中...")
        for page_num, page_images in page_to_images.items():
            # 找到对应的页面
            matching_pages = [p for p in pages if p.get('page_number') == page_num]
            if not matching_pages:
                continue
            
            page = matching_pages[0]
            
            # 构建图片描述文本块
            image_descriptions = []
            for img in page_images:
                desc = img.get('description', '[无描述]')
                img_filename = os.path.basename(img.get('image_path', ''))
                
                # 【改进】明确标注页码和图片
                img_block = f"\n\n【第{page_num}页·图片】({img_filename})\n{desc}\n【/图片】\n"
                image_descriptions.append(img_block)
            
            # 将图片描述插入到页面文本的合适位置
            # 通常在该页面文本的开头或结尾
            if image_descriptions:
                # 优先插入到上下文位置
                context = page_images[0].get('context', '')
                if context and context in page['text']:
                    # 在上下文后插入
                    page['text'] = page['text'].replace(
                        context,
                        context + ''.join(image_descriptions)
                    )
                else:
                    # 否则插入到页面末尾
                    page['text'] += ''.join(image_descriptions)
                
                print(f"   ✅ 第{page_num}页插入了{len(page_images)}张图片的描述")
        
        print(f"✅ 图片描述插入完成！\n")
        return images
    
    def _get_file_hash(self, filename: str) -> str:
        """生成文件名的 MD5 哈希，解决中文路径问题"""
        return hashlib.md5(filename.encode('utf-8')).hexdigest()
    
    def check_libreoffice(self) -> bool:
        """检查LibreOffice是否可用"""
        try:
            # 优先使用配置的路径
            if os.path.exists(LIBREOFFICE_PATH):
                result = subprocess.run(
                    [LIBREOFFICE_PATH, '--version'],
                    capture_output=True,
                    timeout=5
                )
                if result.returncode == 0:
                    print(f"   ✅ 使用配置的LibreOffice: {LIBREOFFICE_PATH}")
                    return True
            
            # 回退到系统PATH中的libreoffice
            result = subprocess.run(
                ['libreoffice', '--version'],
                capture_output=True,
                timeout=5
            )
            if result.returncode == 0:
                print(f"   ✅ 使用系统LibreOffice")
                return True
            return False
        except:
            return False
    
    def convert_to_pdf(self, file_path: str) -> Optional[str]:
        """
        使用LibreOffice将PPTX/DOCX转换为PDF（如果可用）
        
        Returns:
            转换后的PDF文件路径，如果失败或LibreOffice不可用则返回None
        """
        # 检查LibreOffice是否可用
        if not self.check_libreoffice():
            print(f"   ℹ️  LibreOffice不可用，将使用基础解析器")
            return None
        
        try:
            # 创建临时目录
            temp_dir = tempfile.mkdtemp()
            
            # 确定使用哪个LibreOffice
            libreoffice_cmd = LIBREOFFICE_PATH if os.path.exists(LIBREOFFICE_PATH) else 'libreoffice'
            
            # 使用LibreOffice转换
            cmd = [
                libreoffice_cmd,
                '--headless',
                '--convert-to', 'pdf',
                '--outdir', temp_dir,
                file_path
            ]
            
            print(f"   🔄 使用LibreOffice转换为PDF: {os.path.basename(file_path)}")
            result = subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                timeout=60  # 60秒超时
            )
            
            if result.returncode == 0:
                # 查找生成的PDF文件
                pdf_filename = os.path.splitext(os.path.basename(file_path))[0] + '.pdf'
                pdf_path = os.path.join(temp_dir, pdf_filename)
                
                if os.path.exists(pdf_path):
                    print(f"   ✅ 转换成功: {pdf_filename}")
                    return pdf_path
                else:
                    print(f"   ⚠️  未找到转换后的PDF文件")
                    return None
            else:
                print(f"   ⚠️  LibreOffice转换失败，使用基础解析器")
                return None
                
        except subprocess.TimeoutExpired:
            print(f"   ⚠️  转换超时，使用基础解析器")
            return None
        except Exception as e:
            print(f"   ⚠️  转换错误: {e}，使用基础解析器")
            return None

    def load_pdf(self, file_path: str, use_cache: bool = True) -> Tuple[List[Dict], List[Dict]]:
        """加载PDF文件，返回文本内容和图片信息（优先使用MinerU，支持缓存）

        要求：
        1. 优先使用MinerU处理PDF文件
        2. 提取文本内容并格式化，保留页码信息
        3. 提取图片信息
        4. 支持JSON和MD缓存
        
        Args:
            file_path: PDF文件路径
            use_cache: 是否使用缓存
        
        Returns:
            (pages, images): 页面文本列表和图片信息列表
            pages格式: [{"text": str, "page_number": int, "section": str}, ...]
        """
        original_filename = os.path.basename(file_path)
        file_hash = self._get_file_hash(original_filename)
        
        # 定义缓存路径
        json_cache_path = os.path.join(self.save_intermediate_dir, f"{file_hash}.json")
        md_cache_path = os.path.join(self.save_intermediate_dir, f"{file_hash}.md")
        
        # 尝试读取缓存
        if use_cache and os.path.exists(json_cache_path):
            print(f"⚡ [缓存命中] {original_filename}")
            try:
                with open(json_cache_path, 'r', encoding='utf-8') as f:
                    cached_data = json.load(f)
                    # 缓存格式：{"pages": [...], "images": [...]}
                    return cached_data.get('pages', []), cached_data.get('images', [])
            except Exception as e:
                print(f"   ⚠️  缓存读取失败: {e}，重新处理")
        
        pages = []
        images = []
        full_doc_content = []  # 用于生成MD缓存
        
        # 优先使用MinerU
        if self.ocr_processor:
            try:
                # 直接调用 process_file 获取完整结果（包括文本和图片）
                full_result = self.ocr_processor.process_file(file_path)
                
                if full_result and full_result.get('content'):
                    # 【优先使用准确的页码信息】
                    if 'pages' in full_result and full_result['pages']:
                        # 使用从content_list.json提取的准确页码
                        print(f"   ✅ 使用准确的页码信息 ({len(full_result['pages'])} 页)")
                        for page_info in full_result['pages']:
                            page_num = page_info['page_number']
                            page_content = page_info['content']
                            
                            # 【关键修改】将页码明确标注在内容前面
                            text_with_page = f"【第{page_num}页】\n{page_content}\n【/第{page_num}页】"
                            
                            pages.append({
                                "text": text_with_page,
                                "page_number": page_num,
                                "section": ""  # 可以后续从标题提取
                            })
                    else:
                        # 回退到原来的估算方法
                        print(f"   ⚠️  未找到页码信息，使用估算方法")
                        markdown_content = full_result['content']
                        
                        # 按段落分割并估计页码
                        lines = markdown_content.split('\n')
                        current_page = 1
                        current_section = ""
                        current_text = []
                        
                        for line in lines:
                            if line.strip().startswith('#'):
                                if current_text:
                                    text_content = '\n'.join(current_text)
                                    # 【标注页码】
                                    text_with_page = f"【第{current_page}页】\n{text_content}\n【/第{current_page}页】"
                                    pages.append({
                                        "text": text_with_page,
                                        "page_number": current_page,
                                        "section": current_section
                                    })
                                    current_text = []
                                    current_page += 1
                                
                                current_section = line.strip().replace('#', '').strip()
                                current_text.append(line)
                            else:
                                current_text.append(line)
                                
                                if len('\n'.join(current_text)) > 500 and line.strip() == '':
                                    text_content = '\n'.join(current_text)
                                    # 【标注页码】
                                    text_with_page = f"【第{current_page}页】\n{text_content}\n【/第{current_page}页】"
                                    pages.append({
                                        "text": text_with_page,
                                        "page_number": current_page,
                                        "section": current_section
                                    })
                                    current_text = []
                                    current_page += 1
                        
                        if current_text:
                            text_content = '\n'.join(current_text)
                            # 【标注页码】
                            text_with_page = f"【第{current_page}页】\n{text_content}\n【/第{current_page}页】"
                            pages.append({
                                "text": text_with_page,
                                "page_number": current_page,
                                "section": current_section
                            })
                        
                        if not pages:
                            # 【标注页码】
                            text_with_page = f"【第1页】\n{markdown_content}\n【/第1页】"
                            pages.append({
                                "text": text_with_page,
                                "page_number": 1,
                                "section": "全文"
                            })
                    
                    # 提取图片信息
                    if 'images' in full_result and full_result['images']:
                        images = full_result['images']
                        print(f"   📸 提取到 {len(images)} 张图片")
                        
                        # 【新增】为图片生成描述并插入到文本内容中
                        if self.image_describer and images:
                            images = self._add_image_descriptions_and_embed(images, pages)
                    
                    print(f"   📄 分割为 {len(pages)} 个段落")
                    
                    # 保存缓存
                    if self.save_intermediate_dir and pages:
                        try:
                            # 保存 JSON 缓存
                            cache_data = {"pages": pages, "images": images}
                            with open(json_cache_path, 'w', encoding='utf-8') as f:
                                json.dump(cache_data, f, ensure_ascii=False, indent=2)
                            
                            # 保存 MD 缓存
                            if full_result.get('content'):
                                with open(md_cache_path, 'w', encoding='utf-8') as f:
                                    f.write(f"# {original_filename}\n\n")
                                    f.write(full_result['content'])
                            
                            print(f"   💾 结果已缓存至: {json_cache_path}")
                        except Exception as e:
                            print(f"   ⚠️  缓存保存失败: {e}")
                    
                    return pages, images
            except Exception as e:
                print(f"⚠️  MinerU处理失败，尝试基础解析器: {e}")
        
        # 如果MinerU不可用或失败，使用基础解析器（需要手动安装PyPDF2）
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            for page_num, page in enumerate(reader.pages, 1):
                text = page.extract_text()
                if text.strip():
                    # 【统一格式】使用相同的页码标注
                    text_with_page = f"【第{page_num}页】\n{text}\n【/第{page_num}页】"
                    pages.append({
                        "text": text_with_page,
                        "page_number": page_num,
                        "section": f"第 {page_num} 页"
                    })
        except Exception as e:
            print(f"❌ 加载PDF文件失败 {file_path}: {e}")
        
        return pages, images

    def load_pptx(self, file_path: str) -> List[Dict]:
        """加载PPTX文件

        策略（优先级从高到低）：
        1. 如果LibreOffice可用：转换为PDF → 用MinerU处理（OCR + 图片）
        2. 如果LibreOffice不可用或转换失败：使用python-pptx提取文本（仅文本）
        
        注意：python-pptx只能提取文本，无法处理图片和复杂布局
        """
        slides = []
        pdf_path = None
        
        # 策略1: 转换为PDF并用MinerU处理
        if self.ocr_processor:
            try:
                pdf_path = self.convert_to_pdf(file_path)
                if pdf_path:
                    # 用MinerU处理转换后的PDF
                    pages, images = self.load_pdf(pdf_path)
                    # 清理临时PDF
                    if os.path.exists(pdf_path):
                        temp_dir = os.path.dirname(pdf_path)
                        shutil.rmtree(temp_dir, ignore_errors=True)
                    
                    if pages:
                        return pages
            except Exception as e:
                print(f"⚠️  PDF转换+MinerU处理失败: {e}")
                # 清理临时文件
                if pdf_path and os.path.exists(pdf_path):
                    temp_dir = os.path.dirname(pdf_path)
                    shutil.rmtree(temp_dir, ignore_errors=True)
        
        # 策略2: 使用python-pptx基础解析器
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for slide_num, slide in enumerate(prs.slides, 1):
                text_parts = []
                
                # 提取幻灯片中所有文本框的内容
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text_parts.append(shape.text)
                
                # 合并所有文本
                slide_text = "\n".join(text_parts)
                
                if slide_text.strip():
                    formatted_text = f"--- 幻灯片 {slide_num} ---\n{slide_text}\n"
                    slides.append({"text": formatted_text})
        except Exception as e:
            print(f"❌ 加载PPTX文件失败 {file_path}: {e}")
        
        return slides

    def load_docx(self, file_path: str) -> str:
        """加载DOCX文件
        
        策略（优先级从高到低）：
        1. 如果LibreOffice可用：转换为PDF → 用MinerU处理（OCR + 图片）
        2. 如果LibreOffice不可用或转换失败：使用docx2txt提取文本（仅文本）
        
        注意：docx2txt只能提取文本，无法处理图片和复杂布局
        """
        pdf_path = None
        
        # 策略1: 转换为PDF并用MinerU处理
        if self.ocr_processor:
            try:
                pdf_path = self.convert_to_pdf(file_path)
                if pdf_path:
                    # 用MinerU处理转换后的PDF
                    pages, images = self.load_pdf(pdf_path)
                    # 清理临时PDF
                    if os.path.exists(pdf_path):
                        temp_dir = os.path.dirname(pdf_path)
                        shutil.rmtree(temp_dir, ignore_errors=True)
                    
                    if pages:
                        # 合并所有页面的文本
                        return "\n".join([p["text"] for p in pages])
            except Exception as e:
                print(f"⚠️  PDF转换+MinerU处理失败: {e}")
                # 清理临时文件
                if pdf_path and os.path.exists(pdf_path):
                    temp_dir = os.path.dirname(pdf_path)
                    shutil.rmtree(temp_dir, ignore_errors=True)
        
        # 策略2: 使用docx2txt基础解析器
        try:
            import docx2txt
            text = docx2txt.process(file_path)
            return text if text else ""
        except Exception as e:
            print(f"❌ 加载DOCX文件失败 {file_path}: {e}")
            return ""

    def load_txt(self, file_path: str) -> str:
        """加载TXT文件
        要求：
        1. 使用open读取TXT文件（注意使用encoding="utf-8"）
        2. 返回文本内容
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            print(f"加载TXT文件失败 {file_path}: {e}")
            return ""
    
    def load_md(self, file_path: str) -> str:
        """加载Markdown文件
        要求：
        1. 使用open读取MD文件（注意使用encoding="utf-8"）
        2. 返回文本内容
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            print(f"加载Markdown文件失败 {file_path}: {e}")
            return ""

    def load_document(self, file_path: str) -> Tuple[List[Dict[str, str]], List[Dict]]:
        """加载单个文档，PDF和PPT按页/幻灯片分割，返回文档块列表和图片列表
        
        Returns:
            (documents, images): 文档块列表和图片信息列表
        """
        ext = os.path.splitext(file_path)[1].lower()
        filename = os.path.basename(file_path)
        documents = []
        images = []

        if ext == ".pdf":
            pages, images = self.load_pdf(file_path)
            for page_idx, page_data in enumerate(pages, 1):
                documents.append(
                    {
                        "content": page_data["text"],
                        "filename": filename,
                        "filepath": file_path,
                        "filetype": ext,
                        "page_number": page_data.get("page_number", page_idx),
                        "section": page_data.get("section", ""),
                    }
                )
        elif ext == ".pptx":
            slides = self.load_pptx(file_path)
            for slide_idx, slide_data in enumerate(slides, 1):
                documents.append(
                    {
                        "content": slide_data["text"],
                        "filename": filename,
                        "filepath": file_path,
                        "filetype": ext,
                        "page_number": slide_data.get("page_number", slide_idx),
                        "section": slide_data.get("section", f"幻灯片 {slide_idx}"),
                    }
                )
        elif ext == ".docx":
            content = self.load_docx(file_path)
            if content:
                documents.append(
                    {
                        "content": content,
                        "filename": filename,
                        "filepath": file_path,
                        "filetype": ext,
                        "page_number": 1,
                        "section": "文档内容",
                    }
                )
        elif ext == ".txt":
            content = self.load_txt(file_path)
            if content:
                documents.append(
                    {
                        "content": content,
                        "filename": filename,
                        "filepath": file_path,
                        "filetype": ext,
                        "page_number": 1,
                        "section": "文本内容",
                    }
                )
        elif ext == ".md":
            content = self.load_md(file_path)
            if content:
                documents.append(
                    {
                        "content": content,
                        "filename": filename,
                        "filepath": file_path,
                        "filetype": ext,
                        "page_number": 1,
                        "section": "Markdown内容",
                    }
                )
        else:
            print(f"不支持的文件格式: {ext}")

        return documents, images

    def load_all_documents(self) -> Tuple[List[Dict[str, str]], List[Dict]]:
        """加载数据目录下的所有文档
        
        Returns:
            (documents, images): 所有文档块列表和所有图片信息列表
        """
        if not os.path.exists(self.data_dir):
            print(f"数据目录不存在: {self.data_dir}")
            return [], []

        documents = []
        all_images = []

        for root, dirs, files in os.walk(self.data_dir):
            for file in files:
                ext = os.path.splitext(file)[1].lower()
                if ext in self.supported_formats:
                    file_path = os.path.join(root, file)
                    print(f"正在加载: {file_path}")
                    doc_chunks, images = self.load_document(file_path)
                    if doc_chunks:
                        documents.extend(doc_chunks)
                    if images:
                        all_images.extend(images)

        return documents, all_images
