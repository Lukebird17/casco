#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
多模态输入处理器（简化版）
按照新的处理逻辑：
1. 纯文本：文字+图片库检索 → 文本模型
2. 图片：描述图片 → 检索 → 原图+context → 多模态模型
3. 文件：读取内容 → 检索 → 文件+context → 多模态模型
"""

import os
import uuid
import tempfile
from pathlib import Path
from typing import Dict, List, Optional
from PIL import Image

from openai import OpenAI
from config import OPENAI_API_KEY, OPENAI_API_BASE, MULTIMODAL_MODEL_NAME


class MultimodalInputHandler:
    """
    多模态输入处理器（简化版）
    """
    
    def __init__(self):
        """
        初始化处理器
        """
        self.client = OpenAI(api_key=OPENAI_API_KEY, base_url=OPENAI_API_BASE)
        self.multimodal_model = MULTIMODAL_MODEL_NAME
        
        # 临时文件目录
        self.temp_dir = Path(tempfile.gettempdir()) / "rag_agent_temp"
        self.temp_dir.mkdir(exist_ok=True)
    
    def handle_image_input(
        self, 
        image: Image.Image,
        text_query: str
    ) -> Dict:
        """
        处理用户上传的图片
        
        新策略：
        1. 保存图片
        2. 使用多模态模型描述图片
        3. 将描述加入text_query
        4. 返回增强后的query和图片路径
        
        Args:
            image: PIL图片对象
            text_query: 用户的文本问题
            
        Returns:
            {
                "enhanced_query": "增强后的查询（包含图片描述）",
                "image_path": "图片路径（用于传给多模态模型）",
                "image_description": "图片描述"
            }
        """
        # 1. 保存临时图片
        temp_image_path = self._save_temp_image(image)
        print(f"📷 已保存图片: {temp_image_path}")
        
        # 2. 描述图片
        print(f"🔍 使用多模态模型描述图片...")
        try:
            description = self._describe_image(temp_image_path)
            print(f"   ✅ 图片描述: {description[:100]}...")
        except Exception as e:
            print(f"   ⚠️  图片描述失败: {e}")
            description = "无法描述图片"
        
        # 3. 增强查询
        enhanced_query = f"{text_query}\n\n【用户上传的图片描述】\n{description}"
        
        return {
            "enhanced_query": enhanced_query,
            "image_path": temp_image_path,
            "image_description": description
        }
    
    def handle_file_input(
        self,
        file_path: str,
        text_query: str
    ) -> Dict:
        """
        处理用户上传的文件
        
        新策略：
        1. 直接读取文件内容（不OCR）
        2. 返回文件内容和原始query
        3. 在app层面用query检索，然后把文件内容也传给多模态模型
        
        Args:
            file_path: 文件路径
            text_query: 用户的文本问题
            
        Returns:
            {
                "query": "用户原始查询",
                "file_content": "文件内容",
                "file_name": "文件名"
            }
        """
        print(f"📄 读取文件: {file_path}")
        
        # 读取文件内容
        try:
            file_content = self._read_file(file_path)
            print(f"   ✅ 读取成功: {len(file_content)} 字符")
        except Exception as e:
            print(f"   ⚠️  读取失败: {e}")
            return {
                "query": text_query,
                "file_content": None,
                "file_name": os.path.basename(file_path),
                "error": str(e)
            }
        
        return {
            "query": text_query,
            "file_content": file_content,
            "file_name": os.path.basename(file_path)
        }
    
    def _save_temp_image(self, image: Image.Image) -> str:
        """保存临时图片"""
        temp_id = uuid.uuid4().hex[:8]
        temp_path = self.temp_dir / f"query_image_{temp_id}.png"
        image.save(temp_path)
        return str(temp_path)
    
    def _describe_image(self, image_path: str) -> str:
        """
        使用多模态模型描述图片
        
        Args:
            image_path: 图片路径
            
        Returns:
            图片描述文本
        """
        import base64
        
        # 读取图片并转换为base64
        with open(image_path, 'rb') as f:
            image_data = f.read()
        image_base64 = base64.b64encode(image_data).decode()
        
        # 构建消息
        messages = [
            {
                "role": "system",
                "content": "你是一个图片分析助手。请详细描述图片内容，包括：1）主要对象和场景；2）文字内容（如果有）；3）图表或数据（如果有）；4）其他重要细节。"
            },
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": "请详细描述这张图片的内容。"
                    },
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:image/png;base64,{image_base64}"
                        }
                    }
                ]
            }
        ]
        
        # 调用API
        response = self.client.chat.completions.create(
            model=self.multimodal_model,
            messages=messages,
            max_tokens=500,
            temperature=0.3
        )
        
        return response.choices[0].message.content
    
    def _read_file(self, file_path: str) -> str:
        """
        读取文件内容
        
        支持：txt, md, py, pdf, docx, pptx
        对于pdf/docx/pptx，直接使用简单的文本提取，不OCR
        
        Args:
            file_path: 文件路径
            
        Returns:
            文件内容
        """
        ext = Path(file_path).suffix.lower()
        
        # 纯文本文件
        if ext in ['.txt', '.md', '.py', '.java', '.cpp', '.c', '.h', '.js', '.html', '.css', '.json', '.xml', '.yaml', '.yml']:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        
        # PDF
        elif ext == '.pdf':
            return self._read_pdf(file_path)
        
        # Word
        elif ext in ['.docx', '.doc']:
            return self._read_docx(file_path)
        
        # PowerPoint
        elif ext in ['.pptx', '.ppt']:
            return self._read_pptx(file_path)
        
        else:
            # 尝试作为文本读取
            try:
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
            except:
                return f"无法读取文件类型: {ext}"
    
    def _read_pdf(self, file_path: str) -> str:
        """简单读取PDF文本（不OCR）"""
        try:
            import PyPDF2
            text = []
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text.append(page.extract_text())
            return '\n\n'.join(text)
        except ImportError:
            return "需要安装 PyPDF2: pip install PyPDF2"
        except Exception as e:
            return f"PDF读取失败: {e}"
    
    def _read_docx(self, file_path: str) -> str:
        """读取Word文档"""
        try:
            import docx
            doc = docx.Document(file_path)
            text = []
            for para in doc.paragraphs:
                text.append(para.text)
            return '\n'.join(text)
        except ImportError:
            return "需要安装 python-docx: pip install python-docx"
        except Exception as e:
            return f"Word文档读取失败: {e}"
    
    def _read_pptx(self, file_path: str) -> str:
        """读取PowerPoint文档"""
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            text = []
            for slide_num, slide in enumerate(prs.slides, 1):
                text.append(f"--- Slide {slide_num} ---")
                for shape in slide.shapes:
                    if hasattr(shape, 'text'):
                        text.append(shape.text)
            return '\n'.join(text)
        except ImportError:
            return "需要安装 python-pptx: pip install python-pptx"
        except Exception as e:
            return f"PowerPoint文档读取失败: {e}"
